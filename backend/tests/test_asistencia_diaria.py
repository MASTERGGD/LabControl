import datetime
from zoneinfo import ZoneInfo

import pytest

from models.catalogo import CatalogoAlumno, GrupoAcademico, PeriodoEscolar
from models.docencia import AsistenciaDocente, CargaDocente, ClaseDocente
from services.asistencia_diaria import asistencia_diaria
from tests.conftest import auth_headers, get_token


@pytest.fixture
def jornada(db, admin_user, monkeypatch):
    corte = datetime.datetime(2026, 9, 22, 10, 30, tzinfo=ZoneInfo('America/Mexico_City'))
    monkeypatch.setattr('services.asistencia_diaria.now_mx', lambda: corte)
    periodo = PeriodoEscolar(clave='SEP-DIC 2026', es_actual=True, activo=True)
    db.add(periodo); db.flush()
    grupos = [GrupoAcademico(periodo_id=periodo.id, carrera='IA' if i < 4 else 'Software',
                            cuatrimestre=1, grupo=str(i), activo=True) for i in range(6)]
    alumnos = [CatalogoAlumno(matricula=f'DIA-{i}', apellido_paterno='Prueba', apellido_materno='Diaria', nombres=f'Alumno {i}',
                             carrera='IA', cuatrimestre=1, grupo='0', periodo=periodo.clave, activo=True) for i in range(5)]
    db.add_all(grupos + alumnos); db.flush()

    def sesion(grupo, estados=None, estado='CERRADA', inicio='08:00', **extra):
        carga = CargaDocente(docente_id=admin_user.id, periodo_id=periodo.id,
                             grupo_academico_id=grupo.id, tipo_actividad='CLASE', actividad_nombre='Materia',
                             dia_semana=1, hora_inicio=inicio, hora_fin='11:00', estado='ACTIVO', activo=True)
        db.add(carga); db.flush()
        if estados is not None:
            clase = ClaseDocente(carga_docente_id=carga.id, fecha=corte.date(), estado=estado, **extra)
            db.add(clase); db.flush()
            for alumno, valor in estados:
                db.add(AsistenciaDocente(clase_docente_id=clase.id, alumno_id=alumno.id, estado=valor))
        db.flush()
        return carga

    sesion(grupos[0], list(zip(alumnos[:4], ['PRESENTE', 'RETARDO', 'JUSTIFICADA', 'FALTA'])))
    sesion(grupos[0], [(alumnos[0], 'PRESENTE'), (alumnos[3], 'PRESENTE')])
    sesion(grupos[0], [(alumnos[4], 'PRESENTE')], estado='ABIERTA')
    sesion(grupos[1])  # horario iniciado, sin captura
    sesion(grupos[2], inicio='15:00')  # no se debe marcar pendiente todavía
    sesion(grupos[3], [(alumnos[4], 'PRESENTE')], motivo_no_impartida='Suspensión')
    sesion(grupos[4], [(alumnos[0], 'PRESENTE')])  # alumno repetido entre grupos/carreras
    sesion(grupos[5], [(alumnos[4], 'PRESENTE')], estado='CORRECCION')
    db.commit()
    return periodo, grupos, alumnos, sesion


def test_corte_deduplica_y_distingue_cobertura(db, jornada):
    periodo, grupos, _, _ = jornada
    data = asistencia_diaria(db, periodo.id)
    assert data['fecha'] == '2026-09-22'
    assert data['corte'].endswith('-06:00')
    r = data['resumen']
    assert r['asistentes'] == 3  # ni justificación, ni abiertas, ni en corrección
    assert r['alumnos_con_registro'] == 4
    assert (r['a_tiempo'], r['retardos'], r['faltaron'], r['justificados']) == (2, 1, 0, 1)
    assert len(data['detalle']) == 7
    assert {d['estado'] for d in data['detalle']} == {'PRESENTE', 'RETARDO', 'FALTA', 'JUSTIFICADA'}
    assert r['grupos_con_lista'] == 2
    assert r['grupos_sin_lista'] == 2
    assert r['grupos_por_iniciar'] == 1
    assert r['listas_confirmadas'] == 3
    assert r['listas_pendientes'] == 3
    por_id = {g['id']: g for g in data['grupos']}
    assert por_id[grupos[0].id]['asistentes'] == 3
    assert por_id[grupos[3].id]['estado'] == 'SIN_ACTIVIDAD'
    assert por_id[grupos[5].id]['listas_en_captura'] == 1
    assert por_id[grupos[5].id]['faltaron'] == 0
    assert sum(c['asistentes'] for c in data['carreras']) == 4  # total institucional deduplicado = 3


def test_filtro_carrera_y_calendario(db, jornada, monkeypatch):
    periodo, _, _, _ = jornada
    data = asistencia_diaria(db, periodo.id, carrera='Software')
    assert data['resumen']['asistentes'] == 1
    assert len(data['grupos']) == 2
    monkeypatch.setattr('services.asistencia_diaria.estado_fecha_academica', lambda *args: {
        'requiere_asistencia': False, 'permite_iniciar_clase': False, 'motivo': 'Suspensión',
    })
    data = asistencia_diaria(db, periodo.id)
    assert data['nota_calendario'] == 'Suspensión'
    # No inventa pendientes por horario en día suspendido; conserva capturas reales.
    assert data['resumen']['grupos_sin_lista'] == 1
    assert data['resumen']['grupos_por_iniciar'] == 0


def test_serie_deduplica_filtra_y_no_proyecta_despues_del_corte(db, jornada):
    periodo, _, _, _ = jornada
    puntos = asistencia_diaria(db, periodo.id)['asistencia_por_horario']['puntos']
    assert [p['hora'] for p in puntos] == ['08:00', '08:30', '09:00', '09:30', '10:00', '10:30']
    assert all(p['asistentes'] == 3 for p in puntos)
    assert puntos[0]['listas_confirmadas'] == 3
    assert puntos[0]['listas_pendientes'] == 3
    assert puntos[0]['grupos_pendientes'] == 3
    filtrado = asistencia_diaria(db, periodo.id, carrera='Software')['asistencia_por_horario']['puntos']
    assert filtrado[0]['asistentes'] == 1
    assert filtrado[0]['listas_pendientes'] == 1


def test_serie_disminuye_al_terminar_clases_y_respeta_reposiciones(db, jornada):
    periodo, grupos, alumnos, sesion = jornada
    # Al terminar a las 09:00 dejan de contar estos alumnos; siguen listas pendientes.
    for clase in db.query(ClaseDocente).filter(ClaseDocente.estado == 'CERRADA').all():
        clase.carga.hora_fin = '09:00'
    sesion(grupos[0], [(alumnos[4], 'RETARDO')], es_reposicion=True,
           hora_inicio_reposicion='09:30', hora_fin_reposicion='10:00')
    db.commit()
    puntos = {p['hora']: p for p in asistencia_diaria(db, periodo.id)['asistencia_por_horario']['puntos']}
    assert puntos['08:30']['asistentes'] == 3
    assert puntos['09:00']['asistentes'] is None
    assert puntos['09:30']['asistentes'] == 1
    assert puntos['10:00']['asistentes'] is None


def test_serie_cero_confirmado_sin_actividad_y_corte_intermedio(db, jornada, monkeypatch):
    periodo, _, _, _ = jornada
    for carga in db.query(CargaDocente).all():
        carga.hora_fin = '09:00'
    for asistencia in db.query(AsistenciaDocente).all():
        asistencia.estado = 'FALTA'
    db.commit()
    puntos = asistencia_diaria(db, periodo.id)['asistencia_por_horario']['puntos']
    assert puntos[0]['asistentes'] == 0  # lista cerrada sin presentes, sí es cero
    assert puntos[-1]['hora'] == '09:00'
    assert puntos[-1]['asistentes'] == 0  # no hay clases activas
    monkeypatch.setattr('services.asistencia_diaria.now_mx', lambda: datetime.datetime(
        2026, 9, 22, 8, 17, tzinfo=ZoneInfo('America/Mexico_City')))
    puntos = asistencia_diaria(db, periodo.id)['asistencia_por_horario']['puntos']
    assert puntos[-1]['hora'] == '08:17'


def test_endpoint_autorizacion_fecha_y_aislamiento(client, db, admin_user, docente_user, jornada):
    periodo, _, _, _ = jornada
    admin = auth_headers(get_token(client, 'admin@test.com', 'AdminPass123'))
    url = '/reportes-academicos/asistencia-diaria'
    response = client.get(url, params={'periodo_id': periodo.id}, headers=admin)
    assert response.status_code == 200, response.text
    assert response.json()['resumen']['asistentes'] == 3
    assert client.get(url, params={'periodo_id': periodo.id, 'fecha': '2026-09-23'}, headers=admin).status_code == 422
    assert client.get(url, params={'periodo_id': periodo.id, 'carrera': 'Ajena'}, headers=admin).status_code == 422
    assert client.get(url, params={'periodo_id': 999999}, headers=admin).status_code == 404
    otro = PeriodoEscolar(clave='ENE-ABR 2027', es_actual=False, activo=True)
    db.add(otro); db.commit()
    empty = client.get(url, params={'periodo_id': otro.id}, headers=admin).json()
    assert empty['resumen']['asistentes'] == 0
    from dependencies import get_current_user
    client.app.dependency_overrides[get_current_user] = lambda: docente_user
    try:
        assert client.get(url, params={'periodo_id': periodo.id}).status_code == 403
    finally:
        client.app.dependency_overrides.pop(get_current_user, None)


def test_desglose_prioriza_retardo_y_falta_sin_duplicar(db, jornada):
    periodo, grupos, alumnos, sesion = jornada
    sesion(grupos[0], [(alumnos[0], 'RETARDO'), (alumnos[2], 'FALTA'), (alumnos[4], 'FALTA')])
    db.commit()
    data = asistencia_diaria(db, periodo.id)
    r = data['resumen']
    assert (r['a_tiempo'], r['retardos'], r['faltaron'], r['justificados']) == (1, 2, 2, 0)
    assert r['asistentes'] == r['a_tiempo'] + r['retardos']
    assert r['alumnos_con_registro'] == r['asistentes'] + r['faltaron'] + r['justificados']
    assert len([d for d in data['detalle'] if d['alumno_id'] == alumnos[0].id]) == 4


def test_excel_desglose_detalle_filtro_y_autorizacion(client, db, admin_user, docente_user, jornada):
    import io
    import openpyxl
    from dependencies import get_current_user

    periodo, grupos, alumnos, _ = jornada
    alumnos[0].nombres = '=1+1'
    alumnos[0].matricula = '=2+2'
    db.commit()
    url = '/reportes-academicos/asistencia-diaria/excel'
    client.app.dependency_overrides[get_current_user] = lambda: admin_user
    try:
        response = client.get(url, params={'periodo_id': periodo.id})
        assert response.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        assert wb.sheetnames == ['Resumen al corte', 'Detalle por alumno y clase']
        ws = wb.worksheets[0]
        assert ws['A13'].value == 'Carrera'
        assert ws.freeze_panes == 'A14'
        assert [ws.cell(14, i).value for i in range(5, 10)] == [3, 2, 1, 0, 1]
        assert ws['E15'].value is None  # Sin lista no significa falta.
        detalle = wb.worksheets[1]
        assert detalle.max_row == 8
        formulas = [c for row in detalle for c in row if c.value == '=2+2']
        assert formulas and all(c.data_type == 's' for c in formulas)
        filtered = client.get(url, params={'periodo_id': periodo.id, 'carrera': 'Software'})
        assert openpyxl.load_workbook(io.BytesIO(filtered.content)).worksheets[1].max_row == 2
        assert client.get(url, params={'periodo_id': periodo.id, 'fecha': '2026-09-23'}).status_code == 422
        client.app.dependency_overrides[get_current_user] = lambda: docente_user
        assert client.get(url, params={'periodo_id': periodo.id}).status_code == 403
    finally:
        client.app.dependency_overrides.pop(get_current_user, None)
